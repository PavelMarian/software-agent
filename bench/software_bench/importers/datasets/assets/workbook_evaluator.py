from __future__ import annotations

import datetime
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


def _transform(value):
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, datetime.time):
        return str(value)[:-3]
    if isinstance(value, datetime.datetime):
        origin = datetime.datetime(1899, 12, 30)
        delta = value - origin
        return round(delta.days + delta.seconds / 86400.0, 0)
    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


def _cells(spec: str):
    from openpyxl.utils.cell import range_boundaries

    if ":" not in spec:
        return [spec]
    min_col, min_row, max_col, max_row = range_boundaries(spec)
    return [
        f"{openpyxl.utils.get_column_letter(col)}{row}"
        for col in range(min_col, max_col + 1)
        for row in range(min_row, max_row + 1)
    ]


def _compare(gold_path: Path, output_path: Path, positions: str) -> bool:
    if not output_path.is_file():
        return False
    gold = openpyxl.load_workbook(gold_path, data_only=True)
    output = openpyxl.load_workbook(output_path, data_only=True)
    for item in positions.split(","):
        item = item.strip()
        if "!" in item:
            sheet, cells = item.split("!", 1)
            sheet = sheet.strip("'")
        else:
            sheet, cells = gold.sheetnames[0], item
        cells = cells.strip("'")
        if sheet not in output:
            return False
        for cell in _cells(cells):
            left = _transform(gold[sheet][cell].value)
            right = _transform(output[sheet][cell].value)
            if left in (None, "") and right in (None, ""):
                continue
            if type(left) is not type(right) or left != right:
                return False
    return True


def _recalculate(path: Path, target_dir: Path) -> Path:
    """Ask LibreOffice to refresh formula caches before value comparison."""
    completed = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(target_dir),
            str(path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    converted = target_dir / f"{path.stem}.xlsx"
    if completed.returncode != 0 or not converted.is_file():
        raise RuntimeError(f"LibreOffice recalculation failed: {completed.stderr}")
    return converted


def main() -> int:
    metadata = json.loads(Path(".benchmark/metadata.json").read_text(encoding="utf-8"))
    solution = Path("solution.py")
    statuses = {}
    if not solution.is_file():
        print(json.dumps({item["id"]: "FAILED" for item in metadata["cases"]}))
        return 0
    with tempfile.TemporaryDirectory(prefix="spreadsheetbench-") as temp:
        temp_root = Path(temp)
        for case in metadata["cases"]:
            try:
                source = Path("inputs") / case["input"]
                output = temp_root / case["input"].replace("_input.", "_output.")
                shutil.copy2(source, output)
                completed = subprocess.run(
                    [sys.executable, str(solution), str(source), str(output)],
                    capture_output=True,
                    text=True,
                    timeout=600,
                    check=False,
                )
                recalc_dir = temp_root / f"recalculated-{case['id']}"
                recalc_dir.mkdir()
                recalculated = _recalculate(output, recalc_dir)
                passed = completed.returncode == 0 and _compare(
                    Path(".benchmark/gold") / case["answer"],
                    recalculated,
                    case["answer_position"],
                )
            except (OSError, RuntimeError, subprocess.SubprocessError, ValueError):
                passed = False
            statuses[case["id"]] = "PASSED" if passed else "FAILED"
    print(json.dumps(statuses, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
